from __future__ import annotations

import csv
import io
import os
from typing import Any, List, Optional

from ..core.ids import (
    blob_id as compute_blob_id,
    document_id as compute_document_id,
    normalize_text as normalize_text_bytes,
    utc_now_iso,
)
from ..types import Blob, Document
from .errors import ProcessingError
from .text import html_to_text, normalize_text


TEXT_EXTS = {".txt", ".md", ".markdown"}
HTML_EXTS = {".html", ".htm"}
CSV_EXTS = {".csv"}
XLSX_EXTS = {".xlsx"}
PARQUET_EXTS = {".parquet"}
IMG_EXTS = {".png", ".jpg", ".jpeg"}


def detect_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext in TEXT_EXTS:
        return "text"
    if ext in HTML_EXTS:
        return "html"
    if ext in CSV_EXTS:
        return "csv"
    if ext in XLSX_EXTS:
        return "xlsx"
    if ext in PARQUET_EXTS:
        return "parquet"
    if ext in IMG_EXTS:
        return "image"
    return "binary"


def _rows_to_markdown(rows: List[List[str]]) -> str:
    if not rows:
        return ""
    header = rows[0]
    body = rows[1:] if len(rows) > 1 else []
    # Build Markdown table
    out = []
    out.append("| " + " | ".join(header) + " |")
    out.append("| " + " | ".join(["---" for _ in header]) + " |")
    for r in body:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out)


def load_document_from_bytes(
    *,
    source_uri: str,
    filename: str,
    data: bytes,
    processing_cfg: Any | None = None,
) -> Document:
    ptype = detect_type(filename)
    meta = {"filename": os.path.basename(filename), "detected_type": ptype}
    text_out: Optional[str] = None
    blobs: Optional[List[Blob]] = None

    # Defaults from config
    normalize_ws = True
    preserve_md = True
    tables_to_md = True
    ocr_enabled = False
    ocr_lang = "eng"
    if processing_cfg is not None:
        text_cfg = getattr(processing_cfg, "text", None) if not isinstance(processing_cfg, dict) else processing_cfg.get("text")
        if text_cfg is not None:
            normalize_ws = getattr(text_cfg, "normalize_whitespace", normalize_ws) if not isinstance(text_cfg, dict) else text_cfg.get("normalize_whitespace", normalize_ws)
            preserve_md = getattr(text_cfg, "preserve_markdown", preserve_md) if not isinstance(text_cfg, dict) else text_cfg.get("preserve_markdown", preserve_md)
        tables_cfg = getattr(processing_cfg, "tables", None) if not isinstance(processing_cfg, dict) else processing_cfg.get("tables")
        if tables_cfg is not None:
            tables_to_md = getattr(tables_cfg, "to_markdown", tables_to_md) if not isinstance(tables_cfg, dict) else tables_cfg.get("to_markdown", tables_to_md)
        images_cfg = getattr(processing_cfg, "images", None) if not isinstance(processing_cfg, dict) else processing_cfg.get("images")
        if images_cfg is not None:
            ocr_cfg = getattr(images_cfg, "ocr", None) if not isinstance(images_cfg, dict) else images_cfg.get("ocr")
            if ocr_cfg is not None:
                ocr_enabled = getattr(ocr_cfg, "enabled", ocr_enabled) if not isinstance(ocr_cfg, dict) else ocr_cfg.get("enabled", ocr_enabled)
                ocr_lang = getattr(ocr_cfg, "lang", ocr_lang) if not isinstance(ocr_cfg, dict) else ocr_cfg.get("lang", ocr_lang)

    if ptype == "text":
        try:
            text = data.decode("utf-8", errors="replace")
        except Exception as e:  # pragma: no cover
            raise ProcessingError(f"Failed to decode text file: {e}") from e
        if not preserve_md and filename.lower().endswith((".md", ".markdown")):
            # Very light markdown removal: strip common markers
            import re

            text = re.sub(r"^[#>*`\\-]+\\s*", "", text, flags=re.MULTILINE)
        text_out = normalize_text(text, normalize_whitespace=normalize_ws)
    elif ptype == "html":
        try:
            html = data.decode("utf-8", errors="replace")
        except Exception as e:  # pragma: no cover
            raise ProcessingError(f"Failed to decode html: {e}") from e
        text_out = normalize_text(html_to_text(html), normalize_whitespace=normalize_ws)
    elif ptype == "csv":
        f = io.StringIO(data.decode("utf-8", errors="replace"))
        reader = csv.reader(f)
        rows = [list(map(str, r)) for r in reader]
        text_out = _rows_to_markdown(rows) if tables_to_md else "\n".join([",".join(r) for r in rows])
        text_out = normalize_text(text_out, normalize_whitespace=normalize_ws)
    elif ptype == "xlsx":
        try:
            import openpyxl  # type: ignore
        except Exception as e:
            raise ProcessingError("XLSX support requires openpyxl (install extras: '.[excel]')") from e
        # openpyxl can load from a file-like object
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
        text_out = _rows_to_markdown(rows) if tables_to_md else "\n".join([",".join(r) for r in rows])
        text_out = normalize_text(text_out, normalize_whitespace=normalize_ws)
    elif ptype == "parquet":
        # Optional: try pyarrow, otherwise raise
        try:
            import pyarrow.parquet as pq  # type: ignore
            import pyarrow as pa  # noqa: F401
            table = pq.read_table(io.BytesIO(data))
            # convert to CSV-like rows (header + limited rows)
            rows = [list(table.schema.names)]
            for batch in table.to_batches():
                for i in range(batch.num_rows):
                    rows.append([str(batch.column(j)[i].as_py()) for j in range(batch.num_columns)])
                    if len(rows) > 50:
                        break
                if len(rows) > 50:
                    break
            text_out = _rows_to_markdown(rows) if tables_to_md else "\n".join([",".join(r) for r in rows])
            text_out = normalize_text(text_out, normalize_whitespace=normalize_ws)
        except Exception as e:
            raise ProcessingError("Parquet support requires pyarrow (or integrate deltalake)") from e
    elif ptype == "image":
        media_type = "image/png" if filename.lower().endswith(".png") else "image/jpeg"
        blob = Blob(id="blob1", media_type=media_type, data=data, metadata={})
        blobs = [blob]
        if ocr_enabled:
            try:
                import pytesseract  # type: ignore
                from PIL import Image  # type: ignore
                import io as _io

                img = Image.open(_io.BytesIO(data))
                ocr_text = pytesseract.image_to_string(img, lang=ocr_lang or "eng")
                text_out = normalize_text(ocr_text, normalize_whitespace=normalize_ws)
            except Exception as e:  # pragma: no cover - pillow may be missing
                raise ProcessingError("OCR requires pytesseract and Pillow installed (.[ocr])") from e
        else:
            text_out = None
    else:
        # Binary: store as blob metadata only
        blob = Blob(id="blob1", media_type="application/octet-stream", data=data, metadata={})
        blobs = [blob]
        text_out = None

    payload = data if text_out is None else normalize_text_bytes(text_out)
    content_type_map = {
        "text": "text/plain; charset=utf-8",
        "html": "text/html; charset=utf-8",
        "csv": "text/csv; charset=utf-8",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "parquet": "application/x-parquet",
        "image": media_type if 'media_type' in locals() else "image/octet-stream",
    }
    content_type = content_type_map.get(ptype, "application/octet-stream")
    doc_id = compute_document_id(
        source_uri=source_uri,
        payload=payload or b"",
        content_type=content_type,
        content_length=len(payload or b""),
    )
    provenance = {
        "source_uri": source_uri,
        "root_filename": os.path.basename(filename),
        "hash": doc_id.split("_", 1)[-1],
        "created_at": utc_now_iso(),
    }
    if blobs:
        managed: List[Blob] = []
        for blob in blobs:
            data_bytes = blob.data or b""
            managed.append(
                blob.with_id(
                    compute_blob_id(
                        document_id=doc_id,
                        media_type=blob.media_type,
                        payload=data_bytes,
                    )
                )
            )
        blobs = managed

    return Document(id=doc_id, source_uri=source_uri, text=text_out, blobs=blobs, metadata=meta, provenance=provenance)


__all__ = ["detect_type", "load_document_from_bytes"]
